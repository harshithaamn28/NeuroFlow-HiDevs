from docx import Document

from .base import ExtractedPage


class DOCXExtractor:

    async def extract(self, docx_path: str):

        document = Document(docx_path)

        pages = []
        page_number = 1

        # Paragraphs
        for para in document.paragraphs:

            if not para.text.strip():
                continue

            level = None

            if para.style.name.startswith("Heading"):
                parts = para.style.name.split()
                if len(parts) > 1:
                    level = f"h{parts[1]}"

            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content=para.text,
                    content_type="text",
                    metadata={
                        "level": level,
                        "section": para.text if level else None,
                    },
                )
            )

        # Tables
        for table in document.tables:

            rows = []

            for row in table.rows:
                rows.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content="\n".join(rows),
                    content_type="table",
                    metadata={},
                )
            )

        # Headers
        for section in document.sections:

            header = section.header

            for para in header.paragraphs:

                if para.text.strip():

                    pages.append(
                        ExtractedPage(
                            page_number=page_number,
                            content=para.text,
                            content_type="text",
                            metadata={
                                "header": True
                            },
                        )
                    )

        return pages